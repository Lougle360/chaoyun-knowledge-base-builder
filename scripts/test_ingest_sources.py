from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None

BUILD = Path(__file__).with_name("build_kb.py")
INGEST = Path(__file__).with_name("ingest_sources.py")
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8"}


def rows(path: Path) -> list[dict]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


class IngestTests(unittest.TestCase):
    def test_valid_pdf_is_converted(self) -> None:
        # Minimal real PDF with a page, font, content stream and xref table.
        stream = b'BT /F1 12 Tf 72 720 Td (PDF conversion regression sample) Tj ET'
        objects = [
            b'<< /Type /Catalog /Pages 2 0 R >>',
            b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
            b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
            b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream',
        ]
        data, offsets = b'%PDF-1.4\n', [0]
        for i, obj in enumerate(objects, 1):
            offsets.append(len(data))
            data += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
        xref = len(data)
        data += b'xref\n0 6\n0000000000 65535 f \n'
        data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
        data += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / 'input', base / 'output'
            input_root.mkdir()
            (input_root / '有效.pdf').write_bytes(data)
            self.initialize(output_root)
            result = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            source = rows(output_root / 'sources.jsonl')[0]
            self.assertEqual(source['status'], 'ingested')
            self.assertIn('PDF conversion regression sample', (output_root / source['derived_text_path']).read_text(encoding='utf-8'))

    def test_pdf_header_without_structure_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / 'input', base / 'output'
            input_root.mkdir()
            (input_root / '伪造.pdf').write_bytes(b'%PDF-1.4\nnot a real document')
            self.initialize(output_root)
            result = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            self.assertEqual(rows(output_root / 'sources.jsonl')[0]['status'], 'failed')

    def execute(self, *arguments: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, *arguments],
            text=True,
            encoding="utf-8",
            capture_output=True,
            check=False,
            env=ENV,
        )

    def initialize(self, output: Path) -> None:
        result = self.execute(
            str(BUILD),
            "init",
            str(output),
            "--kb-id",
            "demo.kb",
            "--title",
            "演示知识库",
            "--question",
            "如何验证资料摄取",
            "--include",
            "验证",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_text_ingest_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            (input_root / "资料.md").write_text("# 验证\n\n验证资料内容。\n", encoding="utf-8")
            self.initialize(output_root)
            first = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
            source_count = len(rows(output_root / "sources.jsonl"))
            fragment_count = len(rows(output_root / "fragments.jsonl"))
            second = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(second.returncode, 0, second.stdout + second.stderr)
            self.assertEqual(len(rows(output_root / "sources.jsonl")), source_count)
            self.assertEqual(len(rows(output_root / "fragments.jsonl")), fragment_count)

    def test_modified_source_creates_revision(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            source = input_root / "资料.txt"
            source.write_text("验证版本一。\n", encoding="utf-8")
            self.initialize(output_root)
            self.assertEqual(self.execute(str(INGEST), str(input_root), str(output_root)).returncode, 0)
            source.write_text("验证版本二。\n", encoding="utf-8")
            self.assertEqual(self.execute(str(INGEST), str(input_root), str(output_root)).returncode, 0)
            ingested = [row for row in rows(output_root / "sources.jsonl") if row["status"] == "ingested"]
            self.assertEqual(len(ingested), 2)
            self.assertEqual(len({row["revision_id"] for row in ingested}), 2)

    def test_deleted_source_creates_tombstone(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            source = input_root / "资料.txt"
            source.write_text("验证删除。\n", encoding="utf-8")
            self.initialize(output_root)
            self.assertEqual(self.execute(str(INGEST), str(input_root), str(output_root)).returncode, 0)
            source.unlink()
            self.assertEqual(self.execute(str(INGEST), str(input_root), str(output_root)).returncode, 0)
            self.assertEqual(
                sum(row["status"] == "tombstone" for row in rows(output_root / "sources.jsonl")),
                1,
            )
            self.assertEqual(self.execute(str(INGEST), str(input_root), str(output_root)).returncode, 0)
            self.assertEqual(
                sum(row["status"] == "tombstone" for row in rows(output_root / "sources.jsonl")),
                1,
            )

    def test_failed_conversion_is_recorded(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            (input_root / "损坏.pdf").write_bytes(b"not-a-real-pdf")
            self.initialize(output_root)
            result = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(result.returncode, 1)
            self.assertTrue(any(row["status"] == "failed" for row in rows(output_root / "sources.jsonl")))
            report = json.loads(
                (output_root / "validation" / "ingest-report.json").read_text(encoding="utf-8")
            )
            self.assertFalse(report["passed"])

    @unittest.skipUnless(Document is not None, "python-docx 未安装")
    def test_docx_is_converted_with_markitdown(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            document = Document()
            document.add_heading("验证文档", level=1)
            document.add_paragraph("这是跨格式验证内容。")
            document.save(input_root / "资料.docx")
            self.initialize(output_root)
            result = self.execute(str(INGEST), str(input_root), str(output_root))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            source = rows(output_root / "sources.jsonl")[0]
            self.assertEqual(source["parser"]["name"], "markitdown")
            self.assertFalse(source["extensions"]["native_anchor_available"])
            derived = output_root / source["derived_text_path"]
            self.assertIn("跨格式验证内容", derived.read_text(encoding="utf-8"))

    def test_run_defaults_to_three_file_pilot(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root, output_root = base / "input", base / "output"
            input_root.mkdir()
            for index in range(4):
                (input_root / f"资料{index}.md").write_text(
                    f"# 验证\n\n验证资料 {index}。\n",
                    encoding="utf-8",
                )
            result = self.execute(
                str(BUILD),
                "run",
                str(input_root),
                str(output_root),
                "--kb-id",
                "demo.kb",
                "--title",
                "演示知识库",
                "--question",
                "如何验证试点模式",
                "--include",
                "验证",
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(
                len([row for row in rows(output_root / "sources.jsonl") if row["status"] == "ingested"]),
                3,
            )
            self.assertTrue((output_root / "validation" / "review-queue.json").is_file())
            self.assertTrue((output_root / "validation" / "report.json").is_file())


if __name__ == "__main__":
    unittest.main()
